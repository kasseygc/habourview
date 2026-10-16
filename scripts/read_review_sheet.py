"""Read a completed review sheet back into the evaluation sets.

    uv run --with openpyxl python scripts/read_review_sheet.py                 # report only
    uv run --with openpyxl python scripts/read_review_sheet.py --apply         # also update the gold files

Reads docs/review/evaluation_review_completed.xlsx (or the path given). For each row marked N:
- Questions: the JSON after "Corrected spec:" is validated against the catalogue (the same checks a
  model's spec gets). Valid corrections replace the gold spec in eval/questions.yaml.
- Notes: the categories after "Expected categories:" ("(nothing)", or "field: label, ...") are
  checked against the notes schema and replace the expected findings in eval/notes_gold.yaml.
A correction that doesn't validate is reported and left for a person; nothing is guessed.
Every run writes eval/REVIEW.md: what was confirmed, corrected, applied or left, and why.
Rows are matched by ID and the question or note text must still match, so a sheet made from an
older version of the sets can't overwrite the wrong case.
"""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import yaml
from openpyxl import load_workbook

from axon.agent.spec import QuerySpec
from axon.config import ROOT
from axon.notes.schema import FIELDS, label

DEFAULT_SHEET = ROOT / "docs" / "review" / "evaluation_review_completed.xlsx"
QUESTIONS = ROOT / "eval" / "questions.yaml"
NOTES = ROOT / "eval" / "notes_gold.yaml"
REPORT = ROOT / "eval" / "REVIEW.md"


@dataclass
class Row:
    sheet: str
    id: str
    text: str
    ok: str  # Y, N or "" (not reviewed)
    correction: str
    status: str = ""  # confirmed | not reviewed | applied | left for a person | out of date
    change: str = ""
    reason: str = ""
    new: object = field(default=None, repr=False)


def read_sheet(path) -> tuple[list[Row], list[Row]]:
    wb = load_workbook(path, read_only=True, data_only=True)

    def rows(name: str, text_col: int) -> list[Row]:
        out = []
        for r in wb[name].iter_rows(min_row=2, values_only=True):
            if not r or not r[0]:
                continue
            out.append(Row(name, str(r[0]).strip(), str(r[text_col] or "").strip(),
                           str(r[5] or "").strip().upper(), str(r[6] or "").strip()))
        return out

    return rows("Questions", 2), rows("Notes", 3)


# ---------------------------------------------------------------- questions

def corrected_spec(text: str) -> dict:
    """The JSON object after "Corrected spec:"."""
    m = re.search(r"corrected spec\s*:\s*", text, re.IGNORECASE)
    if not m or "{" not in text[m.end():]:
        raise ValueError("no JSON spec after 'Corrected spec:'")
    start = text.index("{", m.end())
    obj, _ = json.JSONDecoder().raw_decode(text[start:])
    return obj


def check_question(row: Row, case: dict | None) -> None:
    if case is None:
        row.status, row.reason = "out of date", "no case with this ID in eval/questions.yaml"
        return
    if case["question"] != row.text:
        row.status, row.reason = "out of date", "the question text differs from eval/questions.yaml"
        return
    if row.ok == "Y":
        row.status = "confirmed"
        return
    if row.ok != "N":
        row.status = "not reviewed"
        return
    if case["kind"] != "answer":
        row.status, row.reason = "left for a person", "corrections to clarify or decline cases are not read automatically"
        return
    try:
        spec = corrected_spec(row.correction)
        QuerySpec.model_validate(spec)
    except (ValueError, json.JSONDecodeError) as e:
        row.status, row.reason = "left for a person", f"the corrected spec doesn't validate: {str(e).splitlines()[0][:200]}"
        return
    row.new = spec
    row.change = f"{flow(compact(case['spec']))}  →  {flow(compact(spec))}"
    row.status = "applied"


def compact(spec: dict) -> dict:
    """The spec as written in questions.yaml: defaults left out, keys in a fixed order."""
    s = QuerySpec.model_validate(spec).model_dump(mode="json")
    out: dict = {"metric": s["metric"]}
    if s["group_by"]:
        out["group_by"] = s["group_by"]
    if s["filters"]:
        out["filters"] = [{"field": f["field"], **({"op": f["op"]} if f["op"] != "in" else {}), "values": f["values"]}
                          for f in s["filters"]]
    tw = {k: v for k, v in s["time_window"].items() if v is not None}
    if tw.get("preset") != "all":
        out["time_window"] = tw
    out["chart"] = s["chart"]
    return out


def flow(obj) -> str:
    """YAML flow style, matching the hand-written files."""
    if isinstance(obj, dict):
        return "{" + ", ".join(f"{k}: {flow(v)}" for k, v in obj.items()) + "}"
    if isinstance(obj, list):
        return "[" + ", ".join(flow(v) for v in obj) + "]"
    if isinstance(obj, str) and not re.fullmatch(r"[A-Za-z0-9_\-. ]+", obj):
        return json.dumps(obj)
    return str(obj)


def apply_questions(rows: list[Row]) -> None:
    lines = QUESTIONS.read_text().splitlines(keepends=True)
    for row in (r for r in rows if r.status == "applied"):
        start = next(i for i, l in enumerate(lines) if l.strip() == f"- id: {row.id}")
        i = next(j for j in range(start + 1, len(lines)) if lines[j].startswith("  spec:") or lines[j].startswith("- id:"))
        if not lines[i].startswith("  spec:"):
            raise RuntimeError(f"{row.id}: no spec line found")
        lines[i] = f"  spec: {flow(compact(row.new))}\n"
    text = "".join(lines)
    cases = {c["id"]: c for c in yaml.safe_load(text)}
    for row in (r for r in rows if r.status == "applied"):  # what was written must read back as intended
        assert compact(cases[row.id]["spec"]) == compact(row.new), row.id
    QUESTIONS.write_text(text)


# ---------------------------------------------------------------- notes

LABEL_TO_VALUE = {(f, label(v).lower()): v for f, values in FIELDS.items() for v in values}


def corrected_findings(text: str) -> list[str]:
    m = re.search(r"expected categories\s*:\s*(.+?)(?:\.\s|\.$|\n|$)", text, re.IGNORECASE)
    if not m:
        raise ValueError("no 'Expected categories:' in the correction")
    body = m.group(1).strip()
    if body.lower().startswith("(nothing)") or body.lower() in {"nothing", "none"}:
        return []
    out = []
    for part in re.split(r",\s*(?=[a-z_]+\s*:)", body):
        f, _, lab = part.partition(":")
        f, lab = f.strip().lower(), lab.strip().lower()
        if f not in FIELDS:
            raise ValueError(f"unknown field {f!r}")
        value = LABEL_TO_VALUE.get((f, lab)) or (lab.replace(" ", "_") if lab.replace(" ", "_") in FIELDS[f] else None)
        if value is None:
            raise ValueError(f"unknown category {lab!r} for {f}")
        out.append(f"{f}:{value}")
    return out


def check_note(row: Row, notes: list[dict]) -> None:
    i = int(row.id[1:]) - 1 if re.fullmatch(r"g\d+", row.id) else -1
    if not 0 <= i < len(notes) or notes[i]["note"] != row.text:
        row.status, row.reason = "out of date", "the note text differs from eval/notes_gold.yaml"
        return
    if row.ok == "Y":
        row.status = "confirmed"
        return
    if row.ok != "N":
        row.status = "not reviewed"
        return
    try:
        new = corrected_findings(row.correction)
    except ValueError as e:
        row.status, row.reason = "left for a person", str(e)
        return
    row.new = new
    row.change = f"[{', '.join(notes[i]['expect'])}]  →  [{', '.join(new)}]"
    row.status = "applied"
    rest = row.correction[row.correction.lower().find("expected categories"):]
    if re.search(r"\balternative", rest, re.IGNORECASE):
        row.reason = "the correction also suggests an alternative; the first option was applied"


def apply_notes(rows: list[Row], notes: list[dict]) -> None:
    lines = NOTES.read_text().splitlines(keepends=True)
    items = [i for i, l in enumerate(lines) if l.startswith("- ")]
    for row in (r for r in rows if r.status == "applied"):
        n = int(row.id[1:]) - 1
        line = lines[items[n]]
        new_line, k = re.subn(r"expect: \[[^\]]*\]", f"expect: [{', '.join(row.new)}]", line)
        if k != 1:
            raise RuntimeError(f"{row.id}: couldn't find the expect list")
        lines[items[n]] = new_line
    text = "".join(lines)
    reread = yaml.safe_load(text)
    for row in (r for r in rows if r.status == "applied"):
        assert reread[int(row.id[1:]) - 1]["expect"] == row.new, row.id
    NOTES.write_text(text)


# ---------------------------------------------------------------- report

def report(path, q_rows: list[Row], n_rows: list[Row], applied: bool) -> str:
    def counts(rows: list[Row]) -> str:
        order = ["confirmed", "applied", "left for a person", "not reviewed", "out of date"]
        names = {"applied": "corrected"}
        return ", ".join(f"{sum(r.status == s for r in rows)} {names.get(s, s)}" for s in order if any(r.status == s for r in rows))

    out = [
        "# Human review of the evaluation sets", "",
        f"Read from `{path.relative_to(ROOT) if path.is_relative_to(ROOT) else path.name}` on {date.today():%d %b %Y}"
        f" by `scripts/read_review_sheet.py`" + ("; corrections written to the gold files." if applied else
                                                  "; report only, gold files unchanged (run with --apply)."), "",
        f"- Questions ({len(q_rows)}): {counts(q_rows)}",
        f"- Notes ({len(n_rows)}): {counts(n_rows)}", "",
    ]
    changed = [r for r in q_rows + n_rows if r.status not in ("confirmed", "not reviewed")]
    if changed:
        out += ["| Sheet | ID | Item | Result | Change or reason |", "|---|---|---|---|---|"]
        for r in changed:
            detail = " ".join(x for x in (f"`{r.change}`" if r.change else "", r.reason) if x)
            status = r.status if applied or r.status != "applied" else "valid, not applied yet"
            out.append(f"| {r.sheet} | {r.id} | {r.text.replace('|', '/')} | {status} | {detail.replace('|', '/')} |")
    skipped = [r.id for r in q_rows + n_rows if r.status == "not reviewed"]
    if skipped:
        out += ["", f"Not reviewed: {', '.join(skipped)}."]
    return "\n".join(out) + "\n"


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("sheet", nargs="?", default=str(DEFAULT_SHEET))
    ap.add_argument("--apply", action="store_true", help="write valid corrections into the gold files")
    args = ap.parse_args()
    path = Path(args.sheet) if Path(args.sheet).is_absolute() else (ROOT / args.sheet).resolve()

    questions = {c["id"]: c for c in yaml.safe_load(QUESTIONS.read_text())}
    notes = yaml.safe_load(NOTES.read_text())
    q_rows, n_rows = read_sheet(path)
    for r in q_rows:
        check_question(r, questions.get(r.id))
    for r in n_rows:
        check_note(r, notes)

    if args.apply:
        apply_questions(q_rows)
        apply_notes(n_rows, notes)
    text = report(path, q_rows, n_rows, args.apply)
    REPORT.write_text(text)
    print(text)


if __name__ == "__main__":
    main()
